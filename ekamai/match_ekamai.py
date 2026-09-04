#!/usr/bin/env python3
"""
ENET vs RC Reconciliation Script

Matches POS sales receipts (data.xlsx) with SCB Bank Statement ENET transactions.
Outputs in the exact structure of data.xlsx with an added 'Match' column:
- If matched: contains Statement Date/Time in DD/MM/YYYY HH:MM:SS format (e.g., 01/01/2026 22:17:00).
- If not matched (unmatched transfer / ambiguous): contains 'REVIEW' highlighted in YELLOW.
- If non-transfer (Cash / Credit Card): contains '-'.
"""

from __future__ import annotations

import argparse
import os
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ----------------------------------------------------------------------
# 1. DATA PREPARATION & NORMALIZATION
# ----------------------------------------------------------------------

def parse_numeric(val: Any) -> Optional[float]:
    """Normalize numeric amount without altering value."""
    if pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip().replace(",", "")
    if val_str in ["", "-", "None", "nan", "NaN"]:
        return None
    try:
        return float(val_str)
    except ValueError:
        return None


def parse_date(val: Any) -> Optional[datetime]:
    """Parse various date/datetime representations into datetime."""
    if pd.isna(val) or val is None:
        return None
    if isinstance(val, (pd.Timestamp, datetime)):
        return pd.to_datetime(val).to_pydatetime()
    val_str = str(val).strip()
    for fmt in [
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
    ]:
        try:
            return datetime.strptime(val_str, fmt)
        except ValueError:
            continue
    return None


def load_statement_data(statement_file: str, sheet_name: str = "Statement") -> pd.DataFrame:
    """Load and normalize bank statement data."""
    if not os.path.exists(statement_file):
        raise FileNotFoundError(f"Statement file not found: {statement_file}")

    df_raw = pd.read_excel(statement_file, sheet_name=sheet_name)
    rows = []

    for idx, row in df_raw.iterrows():
        raw_dt = row.get("Date/Time")
        if pd.isna(raw_dt):
            continue

        dt = parse_date(raw_dt)
        if dt is None:
            continue

        code = str(row.get("Code", "")).strip() if pd.notna(row.get("Code")) else ""
        channel = str(row.get("Channel", "")).strip() if pd.notna(row.get("Channel")) else ""
        deposit = parse_numeric(row.get("Deposit (Credit)"))
        withdrawal = parse_numeric(row.get("Withdrawal (Debit)"))
        desc = str(row.get("Description", "")).strip() if pd.notna(row.get("Description")) else ""
        note = str(row.get("Note", "")).strip() if pd.notna(row.get("Note")) else ""

        direction = (
            "Credit"
            if deposit is not None and deposit > 0
            else ("Debit" if withdrawal is not None and withdrawal > 0 else "Unknown")
        )
        amount = deposit if deposit is not None and deposit > 0 else (withdrawal if withdrawal is not None else 0.0)

        ref_parts = []
        if note and note != "nan":
            ref_parts.append(note)
        if desc and desc != "nan":
            ref_parts.append(desc)
        ref_str = " | ".join(ref_parts) if ref_parts else f"{channel}_{code}"

        rows.append({
            "stmt_id": len(rows),
            "stmt_row": idx + 2,
            "datetime": dt,
            "date_str": dt.strftime("%d/%m/%Y"),
            "time_str": dt.strftime("%H:%M:%S"),
            "formatted_match_str": dt.strftime("%d/%m/%Y %H:%M:%S"),
            "channel": channel,
            "code": code,
            "direction": direction,
            "amount": amount,
            "deposit": deposit or 0.0,
            "withdrawal": withdrawal or 0.0,
            "reference": ref_str,
            "is_enet": (code.upper() == "ENET" or (direction == "Credit" and channel.upper() == "XP")),
        })

    df_stmt = pd.DataFrame(rows)
    if not df_stmt.empty:
        df_stmt = df_stmt.sort_values(by=["datetime", "stmt_id"]).reset_index(drop=True)
    return df_stmt


# ----------------------------------------------------------------------
# 2. RECONCILIATION MATCHING ENGINE
# ----------------------------------------------------------------------

class ReconciliationEngine:
    """
    Executes shift-window detection, anchor-bounded interval sequence alignment,
    and duplicate amount disambiguation.
    """

    def __init__(self, df_stmt: pd.DataFrame, df_rc_raw: pd.DataFrame):
        self.df_stmt = df_stmt.copy()
        self.df_rc_raw = df_rc_raw.copy()

        # Isolate ENET statement items (incoming transfers)
        self.enet_df = self.df_stmt[self.df_stmt["is_enet"]].copy().reset_index(drop=True)
        self.enet_df["enet_key"] = range(len(self.enet_df))

        # Identify transfer column in RC
        self.transfer_col = None
        for c in ["Tranfer", "Transfer", "โอนเงิน", "โอน"]:
            if c in self.df_rc_raw.columns:
                self.transfer_col = c
                break

        date_col = "date" if "date" in self.df_rc_raw.columns else ("Date" if "Date" in self.df_rc_raw.columns else None)
        self.date_col = date_col

        # Build normalized RC list
        rc_rows = []
        for idx, row in self.df_rc_raw.iterrows():
            d_val = row.get(date_col) if date_col else None
            dt = parse_date(d_val)
            date_str = dt.strftime("%d/%m/%Y") if dt else (str(d_val).strip() if pd.notna(d_val) else "")

            t_val = parse_numeric(row.get(self.transfer_col)) if self.transfer_col else None
            is_transfer = t_val is not None and t_val > 0

            rc_rows.append({
                "rc_id": idx,
                "sales_date": dt,
                "date_str": date_str,
                "transfer_amt": t_val if is_transfer else 0.0,
                "is_transfer": is_transfer,
            })

        self.rc_df = pd.DataFrame(rc_rows)
        self.rc_transfers = self.rc_df[self.rc_df["is_transfer"]].copy().reset_index(drop=True)

        # Match tracking: rc_id -> (enet_key, conf, reason, status)
        self.matched_rc: Dict[int, Tuple[int, float, str, str]] = {}
        self.matched_enet: Dict[int, int] = {}
        self.ambiguous_rc: Dict[int, Tuple[List[int], float, str]] = {}

    def _get_shift_window(self, dt: datetime) -> Tuple[datetime, datetime]:
        """Compute the bank statement timestamp window for a sales date (night shift)."""
        start = datetime(dt.year, dt.month, dt.day, 12, 0, 0)
        end = start + timedelta(hours=36)
        return start, end

    def match_all(self):
        """Execute multi-pass matching strategy."""
        # Pass 1: Unique amounts per shift (Anchor points)
        for s_date, group in self.rc_transfers.groupby("date_str", sort=False):
            if group.empty or group.iloc[0]["sales_date"] is None:
                continue
            dt = group.iloc[0]["sales_date"]
            w_start, w_end = self._get_shift_window(dt)

            stmt_cands = self.enet_df[
                (self.enet_df["datetime"] >= w_start) & (self.enet_df["datetime"] <= w_end)
            ]

            rc_counts = group["transfer_amt"].value_counts()
            stmt_counts = stmt_cands["amount"].value_counts()

            for _, rc_row in group.iterrows():
                amt = rc_row["transfer_amt"]
                rc_id = rc_row["rc_id"]
                if rc_counts[amt] == 1 and stmt_counts.get(amt, 0) == 1:
                    matching_stmt = stmt_cands[stmt_cands["amount"] == amt]
                    if len(matching_stmt) == 1:
                        e_key = matching_stmt.iloc[0]["enet_key"]
                        if e_key not in self.matched_enet and rc_id not in self.matched_rc:
                            time_str = matching_stmt.iloc[0]["time_str"]
                            reason = f"Exact unique amount (THB {amt:,.2f}) at {time_str}"
                            self.matched_rc[rc_id] = (e_key, 1.0, reason, "MATCH")
                            self.matched_enet[e_key] = rc_id

        # Pass 2: Interval Anchor Bounded Sequence Alignment for duplicate amounts
        for _ in range(3):
            newly_added = 0
            for s_date, group in self.rc_transfers.groupby("date_str", sort=False):
                if group.empty or group.iloc[0]["sales_date"] is None:
                    continue
                dt = group.iloc[0]["sales_date"]
                w_start, w_end = self._get_shift_window(dt)

                stmt_cands = self.enet_df[
                    (self.enet_df["datetime"] >= w_start) & (self.enet_df["datetime"] <= w_end)
                ]

                group_sorted = group.sort_values("rc_id").reset_index(drop=True)
                anchors = []
                for idx, row in group_sorted.iterrows():
                    rc_id = row["rc_id"]
                    if rc_id in self.matched_rc:
                        e_key = self.matched_rc[rc_id][0]
                        e_dt = self.enet_df.loc[e_key, "datetime"]
                        anchors.append((idx, rc_id, e_key, e_dt))

                full_anchors = [(-1, None, None, w_start)] + anchors + [(len(group_sorted), None, None, w_end)]

                for a_idx in range(len(full_anchors) - 1):
                    left_a = full_anchors[a_idx]
                    right_a = full_anchors[a_idx + 1]

                    rc_slice = group_sorted.iloc[left_a[0] + 1 : right_a[0]]
                    rc_unmatched = rc_slice[~rc_slice["rc_id"].isin(self.matched_rc)]
                    if rc_unmatched.empty:
                        continue

                    stmt_slice = stmt_cands[
                        (stmt_cands["datetime"] >= left_a[3]) & (stmt_cands["datetime"] <= right_a[3])
                    ]
                    stmt_unmatched = stmt_slice[~stmt_slice["enet_key"].isin(self.matched_enet)]

                    for amt, rc_sub in rc_unmatched.groupby("transfer_amt"):
                        stmt_sub = stmt_unmatched[stmt_unmatched["amount"] == amt]
                        if len(rc_sub) == 1 and len(stmt_sub) == 1:
                            r_id = rc_sub.iloc[0]["rc_id"]
                            e_key = stmt_sub.iloc[0]["enet_key"]
                            time_str = stmt_sub.iloc[0]["time_str"]
                            reason = f"Unique amount (THB {amt:,.2f}) within anchor window at {time_str}"
                            self.matched_rc[r_id] = (e_key, 0.95, reason, "MATCH")
                            self.matched_enet[e_key] = r_id
                            newly_added += 1
                        elif len(rc_sub) == len(stmt_sub) and len(rc_sub) > 1:
                            rc_sub_sorted = rc_sub.sort_values("rc_id")
                            stmt_sub_sorted = stmt_sub.sort_values("datetime")
                            for seq_num, ((_, r_item), (_, s_item)) in enumerate(
                                zip(rc_sub_sorted.iterrows(), stmt_sub_sorted.iterrows()), start=1
                            ):
                                r_id = r_item["rc_id"]
                                e_key = s_item["enet_key"]
                                s_time = s_item["time_str"]
                                reason = f"Sequence-aligned match ({seq_num}/{len(rc_sub)} of THB {amt:,.2f}) at {s_time}"
                                self.matched_rc[r_id] = (e_key, 0.90, reason, "MATCH")
                                self.matched_enet[e_key] = r_id
                                newly_added += 1
            if newly_added == 0:
                break

        # Pass 3: Shift-level sequence alignment
        for s_date, group in self.rc_transfers.groupby("date_str", sort=False):
            if group.empty or group.iloc[0]["sales_date"] is None:
                continue
            dt = group.iloc[0]["sales_date"]
            w_start, w_end = self._get_shift_window(dt)

            stmt_cands = self.enet_df[
                (self.enet_df["datetime"] >= w_start) & (self.enet_df["datetime"] <= w_end)
            ]

            unmatched_rc = group[~group["rc_id"].isin(self.matched_rc)]
            unmatched_stmt = stmt_cands[~stmt_cands["enet_key"].isin(self.matched_enet)]

            for amt, rc_sub in unmatched_rc.groupby("transfer_amt"):
                stmt_sub = unmatched_stmt[unmatched_stmt["amount"] == amt]
                if len(rc_sub) == 1 and len(stmt_sub) == 1:
                    r_id = rc_sub.iloc[0]["rc_id"]
                    e_key = stmt_sub.iloc[0]["enet_key"]
                    s_time = stmt_sub.iloc[0]["time_str"]
                    self.matched_rc[r_id] = (e_key, 0.85, f"Shift unique match at {s_time}", "MATCH")
                    self.matched_enet[e_key] = r_id
                elif len(rc_sub) == len(stmt_sub) and len(rc_sub) > 1:
                    rc_sub_sorted = rc_sub.sort_values("rc_id")
                    stmt_sub_sorted = stmt_sub.sort_values("datetime")
                    for seq_num, ((_, r_item), (_, s_item)) in enumerate(
                        zip(rc_sub_sorted.iterrows(), stmt_sub_sorted.iterrows()), start=1
                    ):
                        r_id = r_item["rc_id"]
                        e_key = s_item["enet_key"]
                        s_time = s_item["time_str"]
                        self.matched_rc[r_id] = (e_key, 0.80, f"Shift sequence match at {s_time}", "MATCH")
                        self.matched_enet[e_key] = r_id

        # Pass 4: Unmatched duplicate amounts -> REVIEW_AMBIGUOUS
        for s_date, group in self.rc_transfers.groupby("date_str", sort=False):
            if group.empty or group.iloc[0]["sales_date"] is None:
                continue
            dt = group.iloc[0]["sales_date"]
            w_start, w_end = self._get_shift_window(dt)

            stmt_cands = self.enet_df[
                (self.enet_df["datetime"] >= w_start) & (self.enet_df["datetime"] <= w_end)
            ]

            unmatched_rc = group[~group["rc_id"].isin(self.matched_rc)]
            unmatched_stmt = stmt_cands[~stmt_cands["enet_key"].isin(self.matched_enet)]

            for amt, rc_sub in unmatched_rc.groupby("transfer_amt"):
                stmt_sub = unmatched_stmt[unmatched_stmt["amount"] == amt]
                if len(stmt_sub) > 0 and len(rc_sub) > 0:
                    cand_times = ", ".join(stmt_sub["time_str"].tolist())
                    cand_keys = stmt_sub["enet_key"].tolist()
                    for _, r_item in rc_sub.iterrows():
                        r_id = r_item["rc_id"]
                        reason = f"Ambiguous: {len(rc_sub)} RC vs {len(stmt_sub)} ENET for THB {amt:,.2f}. Stmt candidates: [{cand_times}]"
                        self.ambiguous_rc[r_id] = (cand_keys, 0.50, reason)


# ----------------------------------------------------------------------
# 3. EXPORT TO EXCEL IN data.xlsx STRUCTURE
# ----------------------------------------------------------------------

def export_matched_data_excel(engine: ReconciliationEngine, output_path: str):
    """
    Export DataFrame preserving exact data.xlsx layout, adding 'Match' column:
    - Matched: Statement Date/Time in DD/MM/YYYY HH:MM:SS format
    - Unmatched transfer / Ambiguous: 'REVIEW' with YELLOW highlight
    - Non-transfer: '-'
    """
    df_raw = engine.df_rc_raw.copy()
    
    match_values = []
    is_review = []

    matched_count = 0
    review_count = 0
    unrelated_count = 0

    for idx in range(len(df_raw)):
        rc_info = engine.rc_df.loc[idx]
        if not rc_info["is_transfer"]:
            match_values.append("-")
            is_review.append(False)
            unrelated_count += 1
        elif idx in engine.matched_rc:
            e_key = engine.matched_rc[idx][0]
            stmt_row = engine.enet_df.loc[e_key]
            # DD/MM/YYYY HH:MM:SS format
            formatted_dt_str = stmt_row["formatted_match_str"]
            match_values.append(formatted_dt_str)
            is_review.append(False)
            matched_count += 1
        else:
            match_values.append("REVIEW")
            is_review.append(True)
            review_count += 1

    df_raw["Match"] = match_values

    # Write to Excel with openpyxl formatting
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.views.sheetView[0].showGridLines = True

    # Font and styles
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="000000")
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    normal_font = Font(name=font_family, size=10)
    bold_font = Font(name=font_family, size=10, bold=True)
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # Bright Yellow Highlight
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    columns = list(df_raw.columns)

    # Write Headers
    for c_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Write Rows
    for r_idx, row in df_raw.iterrows():
        excel_row = r_idx + 2
        review_flag = is_review[r_idx]

        for c_idx, col_name in enumerate(columns, start=1):
            val = row[col_name]
            cell = ws.cell(row=excel_row, column=c_idx)

            if pd.isna(val):
                cell.value = None
            else:
                cell.value = val

            cell.font = normal_font
            cell.border = thin_border

            # Center align date and match columns
            if col_name in ["date", "Match"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Highlight REVIEW rows in YELLOW
            if review_flag:
                cell.fill = yellow_fill
                if col_name == "Match":
                    cell.font = bold_font

    # Auto column width
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col[:100])
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_path)
    print(f"\n[DONE] Saved matched sales data to: {output_path}")
    print(f"       - Total Rows          : {len(df_raw):,}")
    print(f"       - Matched (Statement) : {matched_count:,}")
    print(f"       - REVIEW (Yellow)     : {review_count:,}")
    print(f"       - Non-transfer (-)    : {unrelated_count:,}")


# ----------------------------------------------------------------------
# 4. CLI RUNNER
# ----------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Reconcile data.xlsx with statement.xlsx and add Match column in DD/MM/YYYY format.")
    parser.add_argument("-s", "--statement", default="statement.xlsx", help="Path to statement Excel file (default: statement.xlsx)")
    parser.add_argument("-d", "--data", default="data.xlsx", help="Path to RC sales Excel file (default: data.xlsx)")
    parser.add_argument("-o", "--output", default="data_matched.xlsx", help="Path to output Excel file (default: data_matched.xlsx)")
    parser.add_argument("--sheet-statement", default="Statement", help="Sheet name for statement (default: Statement)")

    args = parser.parse_args()

    print(f"[*] Loading bank statement: {args.statement} [Sheet: {args.sheet_statement}]...")
    df_stmt = load_statement_data(args.statement, sheet_name=args.sheet_statement)
    print(f"    Loaded {len(df_stmt)} statement records ({len(df_stmt[df_stmt['is_enet']])} ENET items)")

    print(f"[*] Loading POS sales data: {args.data}...")
    df_rc_raw = pd.read_excel(args.data)
    print(f"    Loaded {len(df_rc_raw)} rows")

    print("[*] Running reconciliation engine...")
    engine = ReconciliationEngine(df_stmt, df_rc_raw)
    engine.match_all()

    print(f"[*] Exporting results to: {args.output}...")
    export_matched_data_excel(engine, args.output)


if __name__ == "__main__":
    main()
