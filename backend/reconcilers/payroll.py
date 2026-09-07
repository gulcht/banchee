#!/usr/bin/env python3
"""
Standard Payroll Reconciliation Script (ภ.ง.ด.1 vs สปส. 1-10)

Extracts and reconciles monthly payroll records for any company:
- SSO (สปส. 1-10 ส่วนที่ 2): ค่าจ้าง, เงินสมทบ
- PND1 (ภ.ง.ด.1 ใบแนบ): จำนวนเงินได้ที่จ่ายในครั้งนี้, จำนวนเงินภาษีที่หัก และนำส่งในครั้งนี้

Produces a standard 12-month reconciliation workbook with 5 rows per employee:
1. เงินเดือน (จาก SSO ค่าจ้าง)
2. ปกส (จาก SSO เงินสมทบ)
3. รายได้อื่น (จำนวนเงินได้ที่จ่ายในครั้งนี้ - เงินเดือน SSO)
4. จำนวนเงินที่จ่าย (จาก PND1 จำนวนเงินได้ที่จ่ายในครั้งนี้)
5. จำนวนเงินภาษีที่หัก (จาก PND1 จำนวนเงินภาษีที่หัก และนำส่งในครั้งนี้)
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Ensure backend root is in sys.path
BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from converters.pnd1 import extract_pnd1_data
from converters.sso import extract_sso_data, parse_period

MONTH_HEADERS = [
    "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
    "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."
]

ACCOUNTING_FORMAT = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'


def normalize_id(raw_id: Any) -> str:
    """Strip all non-digit characters from Tax ID / ID card."""
    if not raw_id:
        return ""
    return re.sub(r"\D", "", str(raw_id).strip())


def format_tax_id(tax_id: str) -> str:
    """Format 13-digit ID as X-XXXX-XXXXX-XX-X."""
    digits = normalize_id(tax_id)
    if len(digits) == 13:
        return f"{digits[0]}-{digits[1:5]}-{digits[5:10]}-{digits[10:12]}-{digits[12]}"
    return tax_id


def extract_month_from_pnd1_date(date_str: str) -> Optional[int]:
    """Parse month number (1..12) from PND1 'วัน เดือน ปี ที่จ่าย' (DD/MM/YYYY)."""
    if not date_str:
        return None
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(date_str).strip())
    if m:
        try:
            return int(m.group(2))
        except ValueError:
            return None
    return None


def extract_month_from_sso_period(period_str: str) -> Optional[int]:
    """Parse month number (1..12) from SSO metadata period string."""
    m_str, _ = parse_period(period_str)
    if m_str:
        try:
            return int(m_str)
        except ValueError:
            return None
    return None


class PayrollReconciler:
    """Standard multi-company payroll reconciler (PND1 vs SSO)."""

    def __init__(
        self,
        director_salary_from_pnd: bool = True,
        company_name: Optional[str] = None,
        company_tax_id: Optional[str] = None,
    ):
        self.director_salary_from_pnd = director_salary_from_pnd
        self.company_name = company_name or ""
        self.company_tax_id = company_tax_id or ""
        self.company_account_no = ""

        # Key: normalized_id -> employee info
        self.employees: Dict[str, Dict[str, Any]] = {}
        # Order of appearance: list of normalized_ids
        self.employee_order: List[str] = []

        # Monthly records: normalized_id -> month (1..12) -> data dict
        self.records: Dict[str, Dict[int, Dict[str, float]]] = defaultdict(lambda: defaultdict(dict))
        self.years: List[str] = []

    def _register_employee(self, norm_id: str, name: str, source: str):
        """Register employee and keep preferred name (SSO preferred over PND1)."""
        if norm_id not in self.employees:
            self.employees[norm_id] = {
                "tax_id": norm_id,
                "name": name,
                "source": source,
            }
            self.employee_order.append(norm_id)
        elif source == "SSO" and name:
            self.employees[norm_id]["name"] = name
            self.employees[norm_id]["source"] = source

    def process_sso_file(self, pdf_source: str | Path | Any):
        """Extract and ingest an SSO 1-10 Part 2 PDF."""
        df, meta = extract_sso_data(pdf_source)
        period = meta.get("period", "")
        month = extract_month_from_sso_period(period)
        if not month:
            print(f"[WARN] Could not determine month from SSO file: period='{period}'")
            return

        _, sso_year = parse_period(period)
        if sso_year and sso_year not in self.years:
            self.years.append(sso_year)

        # Auto-detect company details from SSO metadata if not explicitly provided
        if not self.company_name and meta.get("company"):
            self.company_name = meta["company"]
        if not self.company_account_no and meta.get("account_no"):
            self.company_account_no = meta["account_no"]

        for _, row in df.iterrows():
            norm_id = normalize_id(row["เลขประจำตัวประชาชน"])
            if not norm_id:
                continue
            name = str(row.get("ชื่อ-ชื่อสกุล", "")).strip()
            self._register_employee(norm_id, name, source="SSO")

            wage = float(row.get("ค่าจ้าง", 0.0) or 0.0)
            contrib = float(row.get("เงินสมทบ", 0.0) or 0.0)

            rec = self.records[norm_id][month]
            rec["wage_sso"] = wage
            rec["sso_contrib"] = contrib
            rec["has_sso"] = 1.0

    def process_pnd1_file(self, pdf_source: str | Path | Any):
        """Extract and ingest a PND1 Attachment PDF."""
        df = extract_pnd1_data(pdf_source)
        if df.empty:
            return

        first_date = df["วัน เดือน ปี ที่จ่าย"].dropna().iloc[0] if "วัน เดือน ปี ที่จ่าย" in df.columns else ""
        month = extract_month_from_pnd1_date(str(first_date))
        if not month:
            print(f"[WARN] Could not determine month from PND1 file: date='{first_date}'")
            return

        # Extract year from date string if available
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(first_date).strip())
        if m:
            pnd_year = int(m.group(3))
            if pnd_year > 2400:
                pnd_year -= 543
            pnd_year_str = str(pnd_year)
            if pnd_year_str not in self.years:
                self.years.append(pnd_year_str)

        for _, row in df.iterrows():
            norm_id = normalize_id(row["เลขประจำตัวผู้เสียภาษีอากร"])
            if not norm_id:
                continue
            name = str(row.get("ชื่อผู้มีเงินได้", "")).strip()
            self._register_employee(norm_id, name, source="PND1")

            paid = float(row.get("จำนวนเงินได้ที่จ่ายในครั้งนี้", 0.0) or 0.0)
            tax = float(row.get("จำนวนเงินภาษีที่หัก และนำส่งในครั้งนี้", 0.0) or 0.0)

            rec = self.records[norm_id][month]
            rec["pnd_paid"] = paid
            rec["pnd_tax"] = tax
            rec["has_pnd"] = 1.0

    def scan_directory(self, dir_path: str | Path):
        """Scan directory and process all SSO and PND1 attachment PDFs."""
        path = Path(dir_path)
        if not path.is_dir():
            raise NotADirectoryError(f"Directory not found: {dir_path}")

        pdf_files = sorted(list(path.glob("**/*.pdf")))
        sso_files = []
        pnd1_files = []

        for f in pdf_files:
            fname = f.name.lower()
            if "form" in fname:
                continue
            if "sso" in fname or "สปส" in fname:
                sso_files.append(f)
            elif "p01" in fname or "ภงด" in fname or "attach" in fname:
                pnd1_files.append(f)
            else:
                # Ambiguous: inspect content
                try:
                    import pdfplumber
                    with pdfplumber.open(f) as pdf:
                        txt = pdf.pages[0].extract_text() or ""
                        if "แบบรายการแสดงการส่งเงินสมทบ" in txt or "สปส. 1-10" in txt:
                            sso_files.append(f)
                        elif "ภ.ง.ด.1" in txt and "ใบแนบ" in txt:
                            pnd1_files.append(f)
                except Exception:
                    pass

        print(f"Found {len(sso_files)} SSO PDFs and {len(pnd1_files)} PND1 PDFs in: {dir_path}")
        for f in sso_files:
            self.process_sso_file(f)
        for f in pnd1_files:
            self.process_pnd1_file(f)

    def calculate_employee_month(self, norm_id: str, month: int) -> Tuple[float, float, float, float, float]:
        """
        Calculate the 5 fields for an employee in a given month:
        Returns: (salary, sso_contrib, income_other, paid, tax)
        """
        data = self.records[norm_id].get(month, {})
        has_sso = bool(data.get("has_sso", False))
        has_pnd = bool(data.get("has_pnd", False))

        wage_sso = data.get("wage_sso", 0.0)
        sso_contrib = data.get("sso_contrib", 0.0)
        paid = data.get("pnd_paid", 0.0)
        tax = data.get("pnd_tax", 0.0)

        if has_sso:
            salary = wage_sso
            # Other income is the difference between PND1 paid amount and SSO salary (e.g. OT)
            other = max(0.0, paid - salary) if has_pnd else 0.0
            return salary, sso_contrib, other, paid, tax
        elif has_pnd:
            # Employee without SSO (e.g. Company Director)
            if self.director_salary_from_pnd:
                salary = paid
                other = 0.0
            else:
                salary = 0.0
                other = paid
            return salary, 0.0, other, paid, tax
        else:
            return 0.0, 0.0, 0.0, 0.0, 0.0

    def generate_excel(self, output_target: str | Path | Any):
        """Generate formatted Excel file matching standard reconciliation template."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "full-year"

        # Fonts & Styles
        font_company = Font(name="Angsana New", size=14, bold=True)
        font_header = Font(name="Angsana New", size=14, bold=True)
        font_regular = Font(name="Angsana New", size=14, bold=False)
        font_bold = Font(name="Angsana New", size=14, bold=True)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        thin_border_side = Side(border_style="thin", color="000000")
        thin_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )
        double_bottom_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=Side(border_style="double", color="000000"),
        )

        fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        fill_subtotal = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")

        # Row 1: Company Title
        title_parts = []
        if self.company_name:
            title_parts.append(self.company_name)
        if self.company_tax_id:
            title_parts.append(f"เลขประจำตัวผู้เสียภาษีอากร : {self.company_tax_id}")
        elif self.company_account_no:
            title_parts.append(f"เลขที่บัญชีนายจ้าง : {self.company_account_no}")

        title_text = " ".join(title_parts) if title_parts else "รายงานเปรียบเทียบเงินเดือนและภาษีหัก ณ ที่จ่าย (ภ.ง.ด.1 vs สปส. 1-10)"
        ws.cell(row=1, column=1, value=title_text).font = font_company
        ws.row_dimensions[1].height = 24

        # Row 2: Headers
        headers = ["ลำดับ", "ชื่อ", "เงินได้"] + MONTH_HEADERS + ["รวม"]
        ws.row_dimensions[2].height = 24
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = font_header
            cell.alignment = align_center
            cell.border = thin_border
            cell.fill = fill_header

        current_row = 3
        seq = 1

        salary_rows = []
        sso_rows = []
        other_rows = []
        paid_rows = []
        tax_rows = []

        for norm_id in self.employee_order:
            emp = self.employees[norm_id]
            formatted_id = format_tax_id(norm_id)
            emp_name = emp["name"]

            # Row 1: เงินเดือน
            r_sal = current_row
            salary_rows.append(r_sal)
            c1 = ws.cell(row=r_sal, column=1, value=seq)
            c1.font = font_regular
            c1.alignment = align_center
            c1.border = thin_border

            c2 = ws.cell(row=r_sal, column=2, value=formatted_id)
            c2.font = font_regular
            c2.alignment = align_left
            c2.border = thin_border

            c3 = ws.cell(row=r_sal, column=3, value="เงินเดือน")
            c3.font = font_regular
            c3.alignment = align_left
            c3.border = thin_border

            # Row 2: ปกส
            r_sso = current_row + 1
            sso_rows.append(r_sso)
            c21 = ws.cell(row=r_sso, column=1, value=None)
            c21.border = thin_border

            c22 = ws.cell(row=r_sso, column=2, value=emp_name)
            c22.font = font_regular
            c22.alignment = align_left
            c22.border = thin_border

            c23 = ws.cell(row=r_sso, column=3, value="ปกส")
            c23.font = font_regular
            c23.alignment = align_left
            c23.border = thin_border

            # Row 3: รายได้อื่น
            r_oth = current_row + 2
            other_rows.append(r_oth)
            ws.cell(row=r_oth, column=1, value=None).border = thin_border
            ws.cell(row=r_oth, column=2, value=None).border = thin_border
            c33 = ws.cell(row=r_oth, column=3, value="รายได้อื่น")
            c33.font = font_regular
            c33.alignment = align_left
            c33.border = thin_border

            # Row 4: จำนวนเงินที่จ่าย
            r_paid = current_row + 3
            paid_rows.append(r_paid)
            ws.cell(row=r_paid, column=1, value=None).border = thin_border
            ws.cell(row=r_paid, column=2, value=None).border = thin_border
            c43 = ws.cell(row=r_paid, column=3, value="จำนวนเงินที่จ่าย")
            c43.font = font_regular
            c43.alignment = align_left
            c43.border = thin_border

            # Row 5: จำนวนเงินภาษีที่หัก
            r_tax = current_row + 4
            tax_rows.append(r_tax)
            ws.cell(row=r_tax, column=1, value=None).border = thin_border
            ws.cell(row=r_tax, column=2, value=None).border = thin_border
            c53 = ws.cell(row=r_tax, column=3, value="จำนวนเงินภาษีที่หัก")
            c53.font = font_regular
            c53.alignment = align_left
            c53.border = thin_border

            # Populate months 1..12
            for m in range(1, 13):
                col = 3 + m
                sal_val, sso_val, oth_val, paid_val, tax_val = self.calculate_employee_month(norm_id, m)

                for r_idx, val in [
                    (r_sal, sal_val),
                    (r_sso, sso_val),
                    (r_oth, oth_val),
                    (r_paid, paid_val),
                    (r_tax, tax_val),
                ]:
                    cell = ws.cell(row=r_idx, column=col, value=val)
                    cell.font = font_regular
                    cell.alignment = align_right
                    cell.number_format = ACCOUNTING_FORMAT
                    cell.border = thin_border

            # Populate column 16 (รวม)
            for r_idx in [r_sal, r_sso, r_oth, r_paid, r_tax]:
                cell = ws.cell(row=r_idx, column=16, value=f"=SUM(D{r_idx}:O{r_idx})")
                cell.font = font_bold if r_idx in (r_sal, r_paid) else font_regular
                cell.alignment = align_right
                cell.number_format = ACCOUNTING_FORMAT
                cell.border = thin_border

            current_row += 5
            seq += 1

        # Summary Section
        summary_defs = [
            ("รวม", "เงินเดือน", salary_rows),
            ("", "ปกส", sso_rows),
            ("", "รายได้อื่น", other_rows),
            ("", "จำนวนเงินที่จ่าย", paid_rows),
            ("", "จำนวนเงินภาษีที่หัก", tax_rows),
        ]

        summary_row_map = {}
        for label_col2, label_col3, target_rows in summary_defs:
            r = current_row
            summary_row_map[label_col3] = r
            ws.cell(row=r, column=1, value=None).border = thin_border

            c2 = ws.cell(row=r, column=2, value=label_col2 if label_col2 else None)
            c2.font = font_bold
            c2.alignment = align_center if label_col2 else align_left
            c2.border = thin_border
            c2.fill = fill_subtotal

            c3 = ws.cell(row=r, column=3, value=label_col3)
            c3.font = font_bold
            c3.alignment = align_left
            c3.border = thin_border
            c3.fill = fill_subtotal

            # Sum for each month D..O
            for m in range(1, 13):
                col = 3 + m
                col_letter = get_column_letter(col)
                if target_rows:
                    sum_formula = "=" + "+".join(f"{col_letter}{tr}" for tr in target_rows)
                else:
                    sum_formula = "0.00"
                cell = ws.cell(row=r, column=col, value=sum_formula)
                cell.font = font_bold
                cell.alignment = align_right
                cell.number_format = ACCOUNTING_FORMAT
                cell.border = thin_border
                cell.fill = fill_subtotal

            # Col 16 (รวม)
            cell_tot = ws.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})")
            cell_tot.font = font_bold
            cell_tot.alignment = align_right
            cell_tot.number_format = ACCOUNTING_FORMAT
            cell_tot.border = thin_border
            cell_tot.fill = fill_subtotal

            current_row += 1

        # Validation Cross-Check Row:
        # Check: จำนวนเงินที่จ่าย - (เงินเดือน + รายได้อื่น) == 0
        r_check = current_row
        ws.cell(row=r_check, column=1, value=None).border = double_bottom_border
        c2 = ws.cell(row=r_check, column=2, value="ตรวจสอบความถูกต้อง")
        c2.font = font_bold
        c2.alignment = align_left
        c2.border = double_bottom_border

        c3 = ws.cell(row=r_check, column=3, value="ผลต่าง (จ่าย - (เงินเดือน+รายได้อื่น))")
        c3.font = font_bold
        c3.alignment = align_left
        c3.border = double_bottom_border

        r_pnd_sum = summary_row_map["จำนวนเงินที่จ่าย"]
        r_sal_sum = summary_row_map["เงินเดือน"]
        r_oth_sum = summary_row_map["รายได้อื่น"]

        for m in range(1, 13):
            col = 3 + m
            col_letter = get_column_letter(col)
            formula = f"={col_letter}{r_pnd_sum}-({col_letter}{r_sal_sum}+{col_letter}{r_oth_sum})"
            cell = ws.cell(row=r_check, column=col, value=formula)
            cell.font = font_bold
            cell.alignment = align_right
            cell.number_format = ACCOUNTING_FORMAT
            cell.border = double_bottom_border

        cell_tot = ws.cell(row=r_check, column=16, value=f"=P{r_pnd_sum}-(P{r_sal_sum}+P{r_oth_sum})")
        cell_tot.font = font_bold
        cell_tot.alignment = align_right
        cell_tot.number_format = ACCOUNTING_FORMAT
        cell_tot.border = double_bottom_border

        # Column widths
        column_widths = {
            "A": 6.0,
            "B": 28.0,
            "C": 18.0,
            "P": 16.0,
        }
        for col_letter in [get_column_letter(c) for c in range(4, 16)]:
            column_widths[col_letter] = 13.5

        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Ensure directory exists if saving to path
        if isinstance(output_target, (str, Path)):
            out_path = Path(output_target)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(out_path)
            print(f"[SUCCESS] Reconciled workbook saved to: {out_path}")
        else:
            wb.save(output_target)
        print(f"Total employees reconciled: {len(self.employee_order)}")

    def generate_csv(self, output_target: str | Path | Any):
        """Generate reconciled summary as CSV with UTF-8 BOM encoding."""
        import csv
        headers = ["ลำดับ", "เลขประจำตัว/ชื่อ", "เงินได้"] + MONTH_HEADERS + ["รวม"]
        rows: List[List[Any]] = [headers]

        seq = 1
        sal_sums = [0.0] * 12
        sso_sums = [0.0] * 12
        oth_sums = [0.0] * 12
        paid_sums = [0.0] * 12
        tax_sums = [0.0] * 12

        for norm_id in self.employee_order:
            emp = self.employees[norm_id]
            formatted_id = format_tax_id(norm_id)
            emp_name = emp["name"]

            sal_vals = []
            sso_vals = []
            oth_vals = []
            paid_vals = []
            tax_vals = []

            for m in range(1, 13):
                sal_v, sso_v, oth_v, paid_v, tax_v = self.calculate_employee_month(norm_id, m)
                sal_vals.append(sal_v)
                sso_vals.append(sso_v)
                oth_vals.append(oth_v)
                paid_vals.append(paid_v)
                tax_vals.append(tax_v)

                sal_sums[m - 1] += sal_v
                sso_sums[m - 1] += sso_v
                oth_sums[m - 1] += oth_v
                paid_sums[m - 1] += paid_v
                tax_sums[m - 1] += tax_v

            rows.append([seq, formatted_id, "เงินเดือน"] + sal_vals + [sum(sal_vals)])
            rows.append(["", emp_name, "ปกส"] + sso_vals + [sum(sso_vals)])
            rows.append(["", "", "รายได้อื่น"] + oth_vals + [sum(oth_vals)])
            rows.append(["", "", "จำนวนเงินที่จ่าย"] + paid_vals + [sum(paid_vals)])
            rows.append(["", "", "จำนวนเงินภาษีที่หัก"] + tax_vals + [sum(tax_vals)])
            seq += 1

        # Summary rows
        rows.append(["", "รวม", "เงินเดือน"] + sal_sums + [sum(sal_sums)])
        rows.append(["", "", "ปกส"] + sso_sums + [sum(sso_sums)])
        rows.append(["", "", "รายได้อื่น"] + oth_sums + [sum(oth_sums)])
        rows.append(["", "", "จำนวนเงินที่จ่าย"] + paid_sums + [sum(paid_sums)])
        rows.append(["", "", "จำนวนเงินภาษีที่หัก"] + tax_sums + [sum(tax_sums)])

        # Validation row
        diff_vals = [paid_sums[i] - (sal_sums[i] + oth_sums[i]) for i in range(12)]
        rows.append(["", "ตรวจสอบความถูกต้อง", "ผลต่าง (จ่าย - (เงินเดือน+รายได้อื่น))"] + diff_vals + [sum(diff_vals)])

        if isinstance(output_target, (str, Path)):
            out_path = Path(output_target)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            with open(out_path, mode="w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerows(rows)
        else:
            # File-like object (e.g. io.BytesIO or io.StringIO)
            import io
            if isinstance(output_target, io.BytesIO):
                text_buf = io.StringIO()
                writer = csv.writer(text_buf)
                writer.writerows(rows)
                output_target.write(text_buf.getvalue().encode("utf-8-sig"))
            else:
                writer = csv.writer(output_target)
                writer.writerows(rows)

    def get_export_filename(self, ext: str = "xlsx") -> str:
        """Get export filename formatted with detected year."""
        clean_ext = ext.lstrip(".")
        year_str = self.years[0] if self.years else datetime.now().strftime("%Y")
        return f"payroll-reconciled-{year_str}.{clean_ext}"


def main():
    parser = argparse.ArgumentParser(description="Standard Payroll Reconciliation (ภ.ง.ด.1 vs สปส. 1-10)")
    parser.add_argument(
        "-d", "--data-dir",
        required=True,
        help="Directory containing SSO and PND1 attachment PDFs",
    )
    parser.add_argument(
        "-o", "--output",
        default=None,
        help="Output Excel filepath (default: <data-dir>/payroll_reconciled.xlsx)",
    )
    parser.add_argument(
        "--company",
        default=None,
        help="Company name override (auto-detected from documents if not specified)",
    )
    parser.add_argument(
        "--tax-id",
        default=None,
        help="Company tax ID override",
    )
    parser.add_argument(
        "--no-director-salary",
        action="store_true",
        help="Set director salary row to 0 instead of taking from PND1 paid amount",
    )

    args = parser.parse_args()

    output_path = args.output
    if not output_path:
        output_path = str(Path(args.data_dir) / "payroll_reconciled.xlsx")

    reconciler = PayrollReconciler(
        director_salary_from_pnd=not args.no_director_salary,
        company_name=args.company,
        company_tax_id=args.tax_id,
    )
    reconciler.scan_directory(args.data_dir)
    reconciler.generate_excel(output_path)


if __name__ == "__main__":
    main()
